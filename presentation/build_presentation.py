"""Build the editable, evidence-scoped 12-slide project presentation.

Presentation-only dependency: python-pptx. Run from the project root.
"""
from pathlib import Path
import json
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'presentation'
OUT.mkdir(exist_ok=True)
prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)
prs.core_properties.title = 'Digital Twin of Indian Fisheries Export Cold Chain for Spoilage Risk Prediction'
prs.core_properties.subject = 'A research prototype for comparing simulated cold-chain thermal exposure'
prs.core_properties.author = 'Cold-chain research project'
prs.core_properties.keywords = 'cold chain, fisheries, digital twin, thermal exposure, simulation'
FONT = 'Segoe UI'
BG = 'FAFAF7'
INK = '172F36'
TEAL = '167C88'
MUTED = '5B6C72'
LINE = 'DCE3E1'
PALE = 'E8F1F0'
WHITE = 'FFFFFF'
ORANGE = 'AC653B'

def rgb(c):
    return RGBColor.from_string(c)

def text(s, x, y, w, h, value, size=20, color=INK, bold=False,
         align=PP_ALIGN.LEFT, font=FONT):
    shape = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = font
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = rgb(color)
        p.alignment = align
        p.space_after = Pt(3)
    return shape

def rect(s, x, y, w, h, fill, radius=False, stroke=None):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                             Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    if stroke:
        shp.line.color.rgb = rgb(stroke)
        shp.line.width = Pt(.8)
    else:
        shp.line.fill.background()
    if radius:
        shp.adjustments[0] = .12
    # Explicitly disable the Office theme's inherited shape shadow.
    from pptx.oxml.xmlchemy import OxmlElement
    shp._element.spPr.append(OxmlElement('a:effectLst'))
    return shp

def line(s, x1, y1, x2, y2, color=LINE, width=1):
    shp = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shp.line.color.rgb = rgb(color)
    shp.line.width = Pt(width)
    from pptx.oxml.xmlchemy import OxmlElement
    shp._element.spPr.append(OxmlElement('a:effectLst'))
    return shp

def circle(s, x, y, d, fill):
    shp = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    shp.line.fill.background()
    from pptx.oxml.xmlchemy import OxmlElement
    shp._element.spPr.append(OxmlElement('a:effectLst'))
    return shp

def base(n, section, title=None, subtitle=None, dark=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(INK if dark else BG)
    fg = WHITE if dark else INK
    muted = 'ADBFBE' if dark else MUTED
    text(s,.65,.36,11,.3,section.upper(),10.5,muted,bold=True)
    line(s,.65,6.98,12.68,6.98,'365057' if dark else LINE,.8)
    text(s,.65,7.10,9,.2,'INDIAN FISHERIES  /  COLD-CHAIN RESEARCH',9,muted)
    text(s,12.05,7.07,.62,.25,f'{n:02d} / 12',9.5,muted,align=PP_ALIGN.RIGHT)
    if title:
        text(s,.65,.91,11.95,1.12,title,33,fg,bold=True)
    if subtitle:
        text(s,.68,2.02,11.6,.62,subtitle,18,muted)
    return s

def note(s, body, sources):
    s.notes_slide.notes_text_frame.text = body + '\n\nEvidence / sources:\n' + '\n'.join(sources)

P2 = 'reports/modeling/PHASE_2_FINAL_RESULTS.md'
P3 = 'reports/digital_twin/final_handoff.md'
P4 = 'reports/integration/final_handoff.md'

# 01 — Title. No invented presenter/institution details.
s = base(1,'Research project',dark=True)
rect(s,.68,1.35,.55,.055,'50B1B1')
text(s,.65,1.73,11.95,3.15,
     'Digital Twin of Indian Fisheries\nExport Cold Chain for\nSpoilage Risk Prediction',39,WHITE,bold=True)
text(s,.70,5.03,10.8,.75,
     'Testing where a chilled-fish export journey picks up heat.',24,'C7DCDA')
text(s,.70,6.30,11,.35,'RESEARCH PROTOTYPE  ·  INDIA → EUROPE  ·  OCTOBER 2026',11,'ADBFBE')
note(s,'The title is the original research objective. What has been delivered is a thermal-exposure research prototype, not a calibrated fish-spoilage predictor. The examples are hypothetical configured routes. This distinction will be explained on slide 11.',
     ['COMPLETE_README.MD',P4])

# 02 — Journey.
s = base(2,'The problem','Fish passes through many hands.\nThe cold chain must stay connected.',
         'An illustrative export journey from India to Europe')
stages = [('Landing','and icing'),('Processing','and packing'),('Cold','storage'),('Road','transfer'),('Airport','handling'),('Air','transport'),('Arrival','inspection')]
xs = [1.25 + i*1.79 for i in range(7)]
line(s,xs[0],3.50,xs[-1],3.50,TEAL,2.2)
for i,(a,b) in enumerate(stages):
    circle(s,xs[i]-.10,3.40,.20,TEAL)
    text(s,xs[i]-.65,3.95,1.38,.79,a+'\n'+b,17.5,INK,align=PP_ALIGN.CENTER)
text(s,.72,5.62,11.9,.65,'Each handover can change the time spent in warm conditions.',25,INK)
note(s,'Walk through a shipment in everyday language: landing and icing, packing, chilled holding, road transfer, airport handling, international transit and arrival inspection. These are stages of the configured example, not a reconstructed observed shipment. The system also includes a conditional hypothetical sea stress configuration. MPEDA supports Indian export context but its item, market and port tables are separate marginals; they do not prove this specific route.',
     ['configs/digital_twin/experiment.yaml',P3,'https://mpeda.gov.in/?page_id=438'])

# 03 — Failure mechanisms.
s = base(3,'The problem','Small disruptions can add heat.',
         'Temperature and time both matter across the journey.')
items = [(0.75,2.98,'Cooling stops','A refrigerated vehicle loses cooling.'),
         (7.00,2.98,'Handling takes longer','Fish waits during loading or inspection.'),
         (0.75,4.75,'Doors stay open','Warm air enters a chilled space.'),
         (7.00,4.75,'Setpoints rise','Storage or transport runs warmer.')]
for x,y,a,b in items:
    rect(s,x,y,.07,.87,TEAL)
    text(s,x+.25,y-.04,5.05,.45,a,25,bold=True)
    text(s,x+.25,y+.55,4.9,.7,b,18,MUTED)
note(s,'Warmer conditions and longer exposure can affect fish quality, but the effect depends on species, product state, initial quality and biological response. This project tests thermal consequences of plausible disruptions. It has not measured these failure frequencies on Indian shipments or quantified resulting spoilage/safety outcomes.',[P3,'configs/digital_twin/experiment.yaml'])

# 04 — Research question.
s = base(4,'Research question')
text(s,.75,1.60,11.4,2.55,
     'Where does a disruption add\nthe most temperature exposure\nto an export journey?',37,INK,bold=True)
line(s,.77,4.70,12.15,4.70,LINE,1)
text(s,.78,5.04,2.1,.35,'LONGER-TERM AIM',11,TEAL,bold=True)
text(s,3.0,4.98,9.05,1.02,'Use paired fish-quality measurements to turn exposure into a tested spoilage forecast.',23,INK)
note(s,'The current answerable research question is about modeled thermal exposure and stage sensitivity. The longer-term objective is biological prediction. A sensor model alone cannot learn a new temperature response, and thermal exposure alone cannot be presented as spoilage probability.',[P3,P4])

# 05 — Tangible solution and actual dashboard.
s = base(5,'What we built','A virtual cold chain we can experiment with.')
text(s,.73,2.46,3.70,1.12,'Change the cooling.\nSee the effect.',25,INK,bold=True)
text(s,.75,4.03,3.55,1.60,'A computer model follows the journey and compares normal cooling with disruptions.',21,MUTED)
text(s,.75,6.07,3.55,.44,'Working local dashboard',13,TEAL,bold=True)
screen = ROOT/'figures/dashboard/redesign/02_road_failure_desktop.png'
s.shapes.add_picture(str(screen),Inches(5.00),Inches(2.15),width=Inches(7.40))
note(s,'This is an actual saved screenshot of the implemented Streamlit dashboard, not a mock-up. Users choose route, scenario, temperature source, reference and channel, run a simulation, compare saved scenarios and download results. The simulator is a configured digital-twin prototype, not a live data-connected replica of an observed shipment. The visible product-lag option is assumed, not measured fish-core temperature.',
     ['app.py','reports/dashboard/final_handoff.md','figures/dashboard/redesign/02_road_failure_desktop.png'])

# 06 — The data, without an exhaustive inventory.
s = base(6,'What we built','Real measurements, used within their limits.',
         'The sources play different roles in the system.')
rows = [
 ('Fish sensors','DaFiF and rainbow-trout sensor readings','Study fish measurements; no matched quality labels'),
 ('Trout quality','21 measured microbial-count assays','A small, separate study at nominal 4 °C'),
 ('Export statistics','Official MPEDA records','Indian trade context; not shipment tracking'),
 ('Thermal records','Citrus reefer and cold-room measurements','Nonfish proxy patterns; not seafood route logs')]
for i,(a,b,c) in enumerate(rows):
    y = 2.85+i*.81
    line(s,.75,y+.67,12.45,y+.67,LINE,.8)
    text(s,.77,y,2.45,.4,a,20,INK,bold=True)
    text(s,3.30,y,4.15,.65,b,18,INK)
    text(s,7.80,y,4.55,.58,c,17,MUTED)
note(s,'Phase 1 curated 27,120 DaFiF readings from 63 workbooks, 774 separate sensory criterion scores, 25,200 trout sensor readings and 21 individual trout TVC assays. Sensor and quality cohorts are different fish; repeated readings are not independent fish. MPEDA has 10,788 annual marginal statistics rows. Proxy tables have 31,140 citrus and 1,411,515 cold-room rows. A separate smoked-catfish study supplies 336 long-format measured endpoint rows; it supports only within-study treatment curves, not fresh-fish kinetics. The model branches are deliberately not fused into a biologically invalid target table. DaFiF source counts and 2,298 local session-label conflicts are documented in the data-quality report; use counted workbook rows, not published descriptions, for totals.',
     ['reports/data/data_quality_report.md','reports/data/scientific_sufficiency_report.md',
      'https://data.mendeley.com/datasets/vx4ptwk3pb/1','https://zenodo.org/records/23600663',
      'https://mpeda.gov.in/?page_id=438','https://data.mendeley.com/datasets/mvr7m6mcgp/2',
      'https://zenodo.org/records/15130001'])

# 07 — Flow and plain-language definition.
s = base(7,'What we built','From a route choice to a stage comparison.')
steps = [('Choose','a route'),('Set normal','cooling'),('Introduce','a disruption'),('Calculate','exposure'),('Compare','the stages')]
for i,(a,b) in enumerate(steps):
    x=.78+i*2.51
    text(s,x,2.23,1.8,.43,f'{i+1:02d}',15,TEAL,bold=True)
    text(s,x,2.78,2.1,.90,a+'\n'+b,22,INK,bold=True)
    if i<4: text(s,x+2.10,2.91,.35,.45,'→',22,MUTED)
rect(s,.78,4.03,11.72,2.17,PALE)
text(s,1.04,4.20,10.95,.42,'Example: temperature stays at 5 °C for 2 hours',20,INK)
text(s,1.25,4.88,2.3,.65,'3 °C',35,INK,bold=True)
text(s,1.25,5.62,2.8,.37,'Above 2 °C reference',15,MUTED)
text(s,4.16,4.99,.5,.53,'×',29,MUTED)
text(s,5.17,4.88,2.1,.65,'2 h',35,INK,bold=True)
text(s,5.17,5.62,2.2,.37,'Duration',15,MUTED)
text(s,7.98,4.99,.5,.53,'=',29,MUTED)
text(s,9.06,4.88,3.05,.65,'6 °C·h',35,TEAL,bold=True)
text(s,9.06,5.62,3.05,.37,'Thermal exposure',15,MUTED)
text(s,.82,6.45,11.65,.35,'This is a comparison measure, not a universal spoilage threshold.',15,MUTED)
note(s,'The simulator carries temperature state through stages using an illustrative first-order response: nominal air lag 0.75 h and assumed product lag 1.5 h. The degree-hour indicator integrates max(temperature minus reference, 0) over time. The simple constant-temperature example gives (5 minus 2) °C × 2 hours = 6 °C·h. This is a teaching example, not a measured shipment or a spoilage threshold. References 2 °C and 4 °C are comparisons, not universal legal or spoilage cutoffs. The current model does not translate degree-hours to biological deterioration. The stable Python API is src.integration.execute; the UI delegates to it rather than implementing separate calculations.',[P3,'reports/integration/architecture_and_contracts.md','src/digital_twin/exposure.py'])

# 08 — Scenarios.
s = base(8,'What we tested','What if the journey does not go to plan?',
         'Each disruption is compared with the same route under normal cooling.')
scenarios=[('Normal cooling','The reference journey'),('Road cooling failure','Cooling stops for 2 hours'),('Airport delay','Extra waiting and an open door'),
           ('Warmer storage','Setpoint rises by 3 °C'),('Flight excursion','Setpoint rises for 4 hours'),('Combined failures','Road, airport and flight together')]
for i,(a,b) in enumerate(scenarios):
    x=.78+(i%3)*4.20; y=2.93+(i//3)*1.63
    line(s,x,y,x+3.72,y,LINE,1)
    text(s,x,y+.19,3.75,.46,a,23,INK,bold=True)
    text(s,x,y+.81,3.72,.65,b,18,MUTED)
text(s,.80,6.40,11.7,.25,'Processing and arrival delays were also tested. All routes and disruptions are hypothetical.',12.5,MUTED)
note(s,'The configured air baseline lasts 47 h: landing/icing 3, processing 5, storage 12, road 5, airport 6, flight 12 and destination 4 h. Eight scenarios include baseline, processing delay, storage setpoint, road failure, airport delay/door, flight excursion, destination delay and combined. Combined is 50 h because the airport stage gains 3 h. The sea configuration is a conditional 375 h software stress case with a 14-day voyage and no evidence that chilled fish remains viable. Phase 4 ran 48 route/mode/scenario simulations: 2 routes × 3 modes × 8 scenarios. Only the air synthetic example is foregrounded here to avoid conflating distinct domains.',
     ['configs/digital_twin/experiment.yaml','configs/digital_twin/sea_route.yaml',P4])

# 09 — Honest result chart, sourced from machine-readable outputs.
s = base(9,'What we found','Road cooling failure added the most exposure.',
         'Among the configured single-disruption air-route scenarios')
df=pd.read_parquet(ROOT/'results/integration/scenario_comparison.parquet')
sel=df[(df.route=='air')&(df['mode']=='synthetic')&(df.channel=='air')&(df.reference_c==2.0)]
labels=[('road_failure','Road cooling failure'),('airport_delay_door','Airport delay / door'),
        ('storage_setpoint','Warmer storage'),('flight_excursion','Flight excursion'),
        ('destination_delay','Arrival delay'),('processing_delay','Processing delay')]
values={r.scenario:float(r.delta_degree_hours) for r in sel.itertuples()}
text(s,.80,2.69,7.50,.3,'ADDED EXPOSURE VS NORMAL COOLING  ·  °C·h ABOVE 2 °C',11,MUTED,bold=True)
for i,(key,label) in enumerate(labels):
    y=3.24+i*.43
    text(s,.80,y-.035,2.72,.34,label,16.5,INK)
    value=values[key]
    rect(s,3.58,y+.01,value/60*4.04,.23,TEAL if i==0 else 'BACDCD')
    text(s,3.73+value/60*4.04,y-.045,.75,.35,f'{value:.1f}',17,INK,bold=i==0)
line(s,8.65,3.15,8.65,6.05,LINE,1)
text(s,9.05,3.23,3.05,.99,'55.8',58,TEAL,bold=True)
text(s,9.10,4.39,2.83,.52,'additional °C·h',22,INK)
text(s,9.10,5.08,2.82,.88,'from a 2-hour\nroad cooling failure',18,MUTED)
text(s,.80,6.42,11.70,.31,'Simulated air · Hypothetical 47 h route · Disruption sizes differ, so this is a conditional ranking.',12.5,MUTED)
note(s,'All plotted values are additional simulated-air degree-hours above 2 °C relative to the same synthetic-air route baseline. Road 55.8025, airport 37.3003, storage 24.5202, flight 19.9999, destination 4.4946, processing 2.9993. Baseline total is 15.0455 °C·h; road-failure total is 70.8480. The combined case adds 113.0374 and totals 128.0829 but is excluded from this single-disruption chart. Airport and delay cases have longer durations; differences include added hours. This ranks intervention effects, not exposure physically occurring inside each stage. These indicators are neither spoilage percentages nor causal field estimates.',
     ['results/integration/scenario_comparison.parquet','reports/integration/scenario_experiments.md'])

# 10 — Robustness, explicitly conditional.
s = base(10,'How we checked it','The road result held across most tested assumptions.')
text(s,.80,2.52,5.0,1.22,'197 / 200',58,TEAL,bold=True)
text(s,.85,3.86,5.15,1.67,'When we changed our simulation assumptions 200 times, road cooling failure was the biggest contributor in 197 cases.',21,INK)
rect(s,.86,5.74,4.80,.12,LINE)
rect(s,.86,5.74,4.80*197/200,.12,TEAL)
text(s,.86,6.07,5.20,.43,'Air route · 2 °C comparison reference',15,MUTED)
line(s,6.55,2.65,6.55,6.18,LINE,1)
text(s,7.03,2.67,5.18,.86,'The ranking is not universal.',27,INK,bold=True)
text(s,7.06,3.95,4.92,1.52,'It changed when we used equal-sized disruptions or changed the comparison temperature.',23,MUTED)
text(s,7.06,5.75,4.98,.79,'These are hypothetical assumption tests, not validation against real fish shipments.',15,ORANGE)
note(s,'The paired assumption exercise varies stage durations ×0.85–1.15, setpoints ±0.5 °C, ambient ±3 °C, air response 0.4–1.2 h, assumed product response 0.75–3 h and event duration ×0.75–1.25. Road is the largest single intervention in 197/200 air draws at 2 °C, 193/200 at 4 °C, and 100/100 sea draws for either reference. Those counts are sensitivity within assumed ranges, not confidence levels or real-world probabilities. Under equal 2 h/+3 °C events, airport and flight each add 6.0 °C·h above 2 °C, while storage adds 2.9018. At 4 °C, the dominant standardized stage changes to destination. The engineering suite also checked deterministic replay, stage-to-route conservation, trace reload, Shapley conservation and numerical step-size sensitivity; the full rebuild passed 41 tests. None of this validates fish biology or field performance.',
     ['reports/integration/robustness_and_uncertainty.md','reports/integration/stage_ranking_and_attribution.md','reports/setup/final_handoff.md'])

# 11 — Capability boundary.
s = base(11,'What it means','We can compare heat exposure.\nSpoilage prediction needs more evidence.')
text(s,.80,2.76,5.45,.39,'AVAILABLE NOW',12,TEAL,bold=True)
text(s,.80,3.40,5.37,1.98,'Simulate a configured journey\nCompare disruption effects\nInspect and download results',23,INK)
line(s,6.53,2.75,6.53,6.24,LINE,1)
text(s,7.03,2.76,5.15,.39,'NEEDED FOR SPOILAGE PREDICTION',12,ORANGE,bold=True)
text(s,7.03,3.40,5.06,2.05,'Fish temperature and quality measured together\nIndependent shipment tests',23,INK)
text(s,.82,6.30,11.6,.37,'The planned spoilage classifier has not been trained: specimen-matched labels are missing.',14,MUTED)
note(s,'There is no validated freshness classifier, fish-specific variable-temperature deterioration model, spoilage probability, shelf-life forecast, food-safety decision or regulatory-compliance engine. The planned Gradient Boosting classifier was not trained because the specimen-level labels are unavailable; saying only “unvalidated” could imply a fitted classifier exists. Phase 2 did produce a separate age-only trout TVC linear model (17 development assays, grouped-CV MAE 0.539 log10 CFU/g). Its three-assay temporal check is post-hoc after a candidate fish overlap was discovered, so no untouched final test exists. Smoked-catfish treatment curves and DaFiF unsupervised clusters also remain separate diagnostics. These are not coupled into the route to claim fish spoilage. The API returns prediction null/NOT_CALIBRATED and explicitly refuses unsupported biological targets.',
     [P2,'reports/integration/model_capabilities_and_blockers.md','reports/data/scientific_sufficiency_report.md'])

# 12 — Delivered system and feasible next research.
s = base(12,'Conclusion',dark=True)
text(s,.75,1.18,11.70,1.55,'We built a working\nvirtual cold chain.',37,WHITE,bold=True)
text(s,.79,2.98,11.55,.54,'It identifies where simulated temperature exposure increases.',24,'C7DCDA')
text(s,.79,3.82,11.55,.85,'The next step is connecting those simulations with real fish-quality measurements to predict spoilage.',23,'C7DCDA')
steps=[('01','Measure','Collect paired fish-temperature and quality data.'),
       ('02','Validate','Test on independent Indian export shipments.'),
       ('03','Prepare research paper','Document findings and evidence limits.')]
for i,(num,a,b) in enumerate(steps):
    x=.82+i*4.19
    line(s,x,4.92,x+3.67,4.92,'365057',1)
    text(s,x,5.13,.6,.32,num,12,'7EC4C0',bold=True)
    text(s,x,5.61,3.67,.47,a,20 if i==2 else 24,WHITE,bold=True)
    text(s,x,6.19,3.72,.60,b,15,'ADBFBE')
note(s,'The delivered contribution is an auditable data pipeline, evidence-scoped study diagnostics, configurable thermal simulation, integrated experiment service and actual dashboard. The immediate research upgrade is linked specimen/lot/logger/assay data at multiple product temperatures and independent exporter/route/later-shipment holdouts. A pilot must determine sample size by independent biological lots and outcome rate, not sensor-row count. The dashboard is already delivered; the next task is biological calibration and field evaluation. Publication is a future research step, not an accepted or submitted paper claim. Full information, provenance, metrics, run commands and limitations are in COMPLETE_README.MD.',
     ['COMPLETE_README.MD','configs/integration/field_linkage_schema.yaml','reports/data/blocked_data_requests.md'])

target=ROOT/'Fisheries_Cold_Chain_Project.pptx'
prs.save(target)

# Structural checks: editable slide count, geometry and required speaker notes.
issues=[]
for i,slide in enumerate(prs.slides,1):
    for shape in slide.shapes:
        if shape.left < 0 or shape.top < 0 or shape.left+shape.width > prs.slide_width+2000 or shape.top+shape.height > prs.slide_height+2000:
            issues.append(f'Slide {i}: off-slide shape {shape.name}')
    if 'Evidence / sources:' not in slide.notes_slide.notes_text_frame.text:
        issues.append(f'Slide {i}: missing evidence notes')
assert len(prs.slides)==12
assert not issues, issues
verification={'slides':12,'pptx':str(target.relative_to(ROOT)),
              'speaker_notes_on_every_slide':True,'all_shapes_within_slide':True,
              'results_source':'results/integration/scenario_comparison.parquet',
              'road_added_degree_hours_above_2c':values['road_failure'],
              'style':'16:9, Segoe UI, warm white, dark ink, muted teal'}
(OUT/'build_verification.json').write_text(json.dumps(verification,indent=2),encoding='utf-8')
(OUT/'speaker_notes.md').write_text('\n\n'.join(f'## Slide {i}\n\n'+slide.notes_slide.notes_text_frame.text for i,slide in enumerate(prs.slides,1)),encoding='utf-8')
print(json.dumps(verification,indent=2))
